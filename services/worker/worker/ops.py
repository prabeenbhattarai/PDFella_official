"""
Document operations. Each takes input paths and an output path inside a private
temporary directory. Heavy third-party tools run as subprocesses with timeouts.
"""
from __future__ import annotations

import os
import secrets
import shutil
import subprocess
from pathlib import Path

import pymupdf as fitz
import pikepdf

from .errors import JobError

TIMEOUT = int(os.environ.get("OP_TIMEOUT_SECONDS", "300"))


def _run(cmd: list[str], cwd: str, timeout: int = TIMEOUT, env: dict | None = None) -> None:
    try:
        subprocess.run(cmd, cwd=cwd, check=True, timeout=timeout, capture_output=True, env={**os.environ, **(env or {})})
    except subprocess.TimeoutExpired as e:
        raise JobError("timeout") from e
    except subprocess.CalledProcessError as e:
        # stderr may contain document text for some tools; log only the exit code.
        raise JobError("failed", f"{cmd[0]} exited with {e.returncode}") from e
    except FileNotFoundError as e:
        raise JobError("unsupported", f"{cmd[0]} is not installed") from e


def _open_pdf(path: str) -> fitz.Document:
    try:
        doc = fitz.open(path)
    except Exception as e:  # noqa: BLE001
        raise JobError("damaged") from e
    if doc.needs_pass:
        raise JobError("encrypted")
    return doc


# ───────────────────────── OCR

def ocr(inputs: list[str], out: str, params: dict, tmp: str) -> None:
    lang = params.get("language", "eng")
    _run(["ocrmypdf", "--skip-text", "--rotate-pages", "--deskew", "--output-type", "pdf", "--jobs", "2", "-l", lang, inputs[0], out], tmp, timeout=max(TIMEOUT, 900))


# ───────────────────────── PDF → Office / HTML

def pdf_to_docx(inputs: list[str], out: str, params: dict, tmp: str) -> None:
    src = inputs[0]
    if params.get("ocr"):
        ocred = str(Path(tmp) / "ocr.pdf")
        ocr([src], ocred, params, tmp)
        src = ocred
    from pdf2docx import Converter

    _open_pdf(src).close()
    cv = Converter(src)
    try:
        cv.convert(out, multi_processing=False)
    finally:
        cv.close()


def pdf_to_xlsx(inputs: list[str], out: str, params: dict, tmp: str) -> None:
    import pdfplumber
    from openpyxl import Workbook

    def cell(v: str | None):
        if v is None:
            return None
        s = v.strip()
        try:
            return float(s.replace(",", "")) if s and s.replace(",", "").replace(".", "", 1).lstrip("-").isdigit() else s
        except ValueError:
            return s

    wb = Workbook()
    wb.remove(wb.active)
    with pdfplumber.open(inputs[0]) as pdf:
        for i, page in enumerate(pdf.pages, 1):
            ws = wb.create_sheet(f"Page {i}")
            tables = page.extract_tables()
            row = 1
            if tables:
                for t in tables:
                    for r in t:
                        for c, v in enumerate(r, 1):
                            ws.cell(row=row, column=c, value=cell(v))
                        row += 1
                    row += 1
            else:
                # No ruled table: fall back to text lines split on runs of spaces.
                for line in (page.extract_text() or "").splitlines():
                    for c, v in enumerate([p for p in line.split("  ") if p.strip()], 1):
                        ws.cell(row=row, column=c, value=cell(v))
                    row += 1
    if not wb.sheetnames:
        wb.create_sheet("Empty")
    wb.save(out)


def pdf_to_pptx(inputs: list[str], out: str, params: dict, tmp: str) -> None:
    from pptx import Presentation
    from pptx.util import Emu

    doc = _open_pdf(inputs[0])
    prs = Presentation()
    first = doc[0].rect
    prs.slide_width = Emu(int(first.width * 12700))
    prs.slide_height = Emu(int(first.height * 12700))
    blank = prs.slide_layouts[6]
    for i, page in enumerate(doc):
        pix = page.get_pixmap(dpi=150)
        img = Path(tmp) / f"p{i}.png"
        pix.save(img)
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(img), 0, 0, width=prs.slide_width, height=prs.slide_height)
        slide.notes_slide.notes_text_frame.text = page.get_text("text")[:20000]
    prs.save(out)


def pdf_to_html(inputs: list[str], out: str, params: dict, tmp: str) -> None:
    doc = _open_pdf(inputs[0])
    parts = ["<!doctype html><meta charset='utf-8'><title>Document</title><body style='background:#eee'>"]
    for page in doc:
        parts.append(f"<div style='margin:16px auto;background:#fff;width:{page.rect.width}pt'>{page.get_text('html')}</div>")
    parts.append("</body>")
    Path(out).write_text("".join(parts), encoding="utf-8")


# ───────────────────────── Office → PDF

def office_to_pdf(inputs: list[str], out: str, params: dict, tmp: str) -> None:
    profile = Path(tmp) / "lo-profile"
    outdir = Path(tmp) / "lo-out"
    outdir.mkdir()
    _run(["soffice", f"-env:UserInstallation=file://{profile}", "--headless", "--norestore", "--convert-to", "pdf", "--outdir", str(outdir), inputs[0]], tmp, timeout=max(TIMEOUT, 300))
    produced = list(outdir.glob("*.pdf"))
    if not produced:
        raise JobError("failed", "no output")
    shutil.move(str(produced[0]), out)


# ───────────────────────── Security

def protect(inputs: list[str], out: str, params: dict, tmp: str) -> None:
    try:
        pdf = pikepdf.open(inputs[0])
    except pikepdf.PasswordError as e:
        raise JobError("encrypted") from e
    except pikepdf.PdfError as e:
        raise JobError("damaged") from e
    perms = pikepdf.Permissions(
        print_lowres=bool(params.get("allowPrint", True)),
        print_highres=bool(params.get("allowPrint", True)),
        extract=bool(params.get("allowCopy", False)),
        accessibility=True,
        modify_annotation=False,
        modify_form=False,
        modify_other=False,
        modify_assembly=False,
    )
    pdf.save(out, encryption=pikepdf.Encryption(user=params["password"], owner=secrets.token_urlsafe(24), R=6, allow=perms))


def unlock(inputs: list[str], out: str, params: dict, tmp: str) -> None:
    try:
        probe = pikepdf.open(inputs[0], password="")
        encrypted = probe.is_encrypted
        probe.close()
        if not encrypted:
            raise JobError("not_encrypted")
    except pikepdf.PasswordError:
        pass
    try:
        pdf = pikepdf.open(inputs[0], password=params["password"])
    except pikepdf.PasswordError as e:
        raise JobError("wrong_password") from e
    pdf.save(out)


def repair(inputs: list[str], out: str, params: dict, tmp: str) -> None:
    try:
        # qpdf (via pikepdf) reconstructs damaged cross-reference tables on open.
        pdf = pikepdf.open(inputs[0], attempt_recovery=True)
        pdf.save(out, fix_metadata_version=True, object_stream_mode=pikepdf.ObjectStreamMode.generate)
        return
    except pikepdf.PasswordError as e:
        raise JobError("encrypted") from e
    except Exception:  # noqa: BLE001
        pass
    try:
        doc = fitz.open(inputs[0])
        doc.save(out, garbage=4, deflate=True, clean=True)
    except Exception as e:  # noqa: BLE001
        raise JobError("damaged") from e


# ───────────────────────── Optimise

PRESETS = {
    "recommended": {"gs": "/ebook", "dpi": 150, "quality": 72},
    "high": {"gs": "/screen", "dpi": 96, "quality": 55},
    "quality": {"gs": "/printer", "dpi": 220, "quality": 85},
}


def compress(inputs: list[str], out: str, params: dict, tmp: str) -> None:
    p = PRESETS.get(params.get("preset", "recommended"), PRESETS["recommended"])
    if shutil.which("gs"):
        _run(["gs", "-sDEVICE=pdfwrite", "-dCompatibilityLevel=1.6", f"-dPDFSETTINGS={p['gs']}", "-dNOPAUSE", "-dQUIET", "-dBATCH", "-dSAFER",
              "-dDetectDuplicateImages=true", f"-sOutputFile={out}", inputs[0]], tmp)
    else:
        doc = _open_pdf(inputs[0])
        if hasattr(doc, "rewrite_images"):
            doc.rewrite_images(dpi_threshold=p["dpi"] + 20, dpi_target=p["dpi"], quality=p["quality"])
        doc.save(out, garbage=4, deflate=True, deflate_images=True, deflate_fonts=True, use_objstms=1)
    if os.path.getsize(out) >= os.path.getsize(inputs[0]):
        shutil.copyfile(inputs[0], out)


# ───────────────────────── Redaction

def _to_fitz_rect(page: fitz.Page, rect: list[float]) -> fitz.Rect:
    """PDF user space (bottom-left origin) → MuPDF unrotated page space (top-left origin)."""
    x, y, w, h = rect
    r = fitz.Rect(x, y, x + w, y + h) * page.transformation_matrix
    r.normalize()
    return r


def redact(inputs: list[str], out: str, params: dict, tmp: str, fill: bool = True) -> None:
    doc = _open_pdf(inputs[0])
    by_page: dict[int, list[dict]] = {}
    for r in params.get("regions", []):
        by_page.setdefault(r["pageIndex"], []).append(r)
    for idx, regs in by_page.items():
        if idx < 0 or idx >= doc.page_count:
            continue
        page = doc[idx]
        for r in regs:
            color = None
            if fill:
                hexv = r.get("fill", "#000000").lstrip("#")
                color = tuple(int(hexv[i:i + 2], 16) / 255 for i in (0, 2, 4))
            page.add_redact_annot(_to_fitz_rect(page, r["rect"]), fill=color)
        # Removes text, vector graphics and the covered image pixels underneath.
        page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_PIXELS, graphics=fitz.PDF_REDACT_LINE_ART_REMOVE_IF_TOUCHED if fill else 0)
    # garbage=4 drops now-unreferenced objects so nothing redacted survives in the file.
    doc.save(out, garbage=4, deflate=True, clean=True)


def edit_text(inputs: list[str], out: str, params: dict, tmp: str) -> None:
    """Remove original text in the given regions without painting a box (replacement is drawn client-side)."""
    redact(inputs, out, params, tmp, fill=False)


OPS = {
    "ocr": ocr,
    "pdf-to-docx": pdf_to_docx,
    "pdf-to-xlsx": pdf_to_xlsx,
    "pdf-to-pptx": pdf_to_pptx,
    "pdf-to-html": pdf_to_html,
    "office-to-pdf": office_to_pdf,
    "protect": protect,
    "unlock": unlock,
    "repair": repair,
    "compress": compress,
    "redact": redact,
    "edit-text": edit_text,
}
